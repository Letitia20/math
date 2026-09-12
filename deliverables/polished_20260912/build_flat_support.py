"""Build and audit the single-level supporting-materials package.

The four official result workbooks are copied byte-for-byte.  They are opened
only for read-only structural and content checks; this script never saves them
through openpyxl.
"""
from __future__ import annotations

import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import zipfile

from openpyxl import load_workbook


OUT = Path(__file__).resolve().parent
ROOT = OUT.parents[1]
SOURCE = ROOT / "deliverables" / "supporting_materials"
TARGET = OUT / "supporting_materials"
STAGING = OUT / "_supporting_materials_flat"
BACKUP = OUT / "_supporting_materials_nested_backup"
TEMPLATE_DIR = (
    Path.home()
    / "Desktop"
    / "math"
    / "72c3104acf279568b3c1891328fa3c7e_1410570938483445441_m"
    / "A题"
    / "附件"
    / "附件3"
)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def style_signature(cell) -> tuple:
    style = cell._style
    return tuple(style) if style is not None else ()


def dimension_signature(dimension) -> dict[str, object]:
    return {
        "width": getattr(dimension, "width", None),
        "height": getattr(dimension, "height", None),
        "hidden": dimension.hidden,
        "outlineLevel": dimension.outlineLevel,
        "collapsed": dimension.collapsed,
    }


def workbook_audit(template: Path, result: Path) -> dict[str, object]:
    template_book = load_workbook(template, data_only=False, read_only=False)
    result_book = load_workbook(result, data_only=False, read_only=False)
    assert template_book.sheetnames == result_book.sheetnames
    sheets = []
    total_cells = 0
    nonempty_cells = 0
    for template_sheet, result_sheet in zip(
        template_book.worksheets, result_book.worksheets, strict=True
    ):
        layout_comparison = {
            "merged_cells_match": set(template_sheet.merged_cells.ranges)
            == set(result_sheet.merged_cells.ranges),
            "freeze_panes_match": template_sheet.freeze_panes == result_sheet.freeze_panes,
            "gridlines_match": template_sheet.sheet_view.showGridLines
            == result_sheet.sheet_view.showGridLines,
            "first_row_dimension_match": dimension_signature(template_sheet.row_dimensions[1])
            == dimension_signature(result_sheet.row_dimensions[1]),
            "first_column_dimension_match": dimension_signature(template_sheet.column_dimensions["A"])
            == dimension_signature(result_sheet.column_dimensions["A"]),
        }

        layout_comparison["header_anchor_style_match"] = (
            style_signature(template_sheet["A1"]) == style_signature(result_sheet["A1"])
        )
        layout_comparison["value_anchor_style_match"] = (
            style_signature(template_sheet["B2"]) == style_signature(result_sheet["B2"])
        )

        values = []
        for row in result_sheet.iter_rows():
            for cell in row:
                total_cells += 1
                if cell.value is not None:
                    nonempty_cells += 1
                if isinstance(cell.value, str):
                    assert not cell.value.startswith("#"), f"Excel error in {result.name}:{cell.coordinate}"
                values.append(cell.value)
        assert result_sheet.max_row > template_sheet.max_row
        assert result_sheet.max_column >= template_sheet.max_column
        assert result_sheet.cell(2, 1).value is not None
        assert result_sheet.cell(result_sheet.max_row, 1).value is not None
        # In result4, positions outside the shrinking body are intentionally
        # blank, but the centre and actual-surface columns must always be filled.
        if result.name == "result4.xlsx":
            assert all(result_sheet.cell(row, 2).value is not None for row in range(2, result_sheet.max_row + 1))
            assert all(
                result_sheet.cell(row, result_sheet.max_column).value is not None
                for row in range(2, result_sheet.max_row + 1)
            )
        else:
            assert all(value is not None for value in values)
        sheets.append(
            {
                "name": result_sheet.title,
                "rows": result_sheet.max_row,
                "columns": result_sheet.max_column,
                "nonempty_cells": sum(
                    cell.value is not None for row in result_sheet.iter_rows() for cell in row
                ),
                "format_preservation": "packaged workbook is byte-identical to the filled source workbook",
                "layout_comparison": layout_comparison,
            }
        )
    template_book.close()
    result_book.close()
    return {
        "filename": result.name,
        "sha256": sha256(result),
        "sheets": sheets,
        "cells_checked": total_cells,
        "nonempty_cells": nonempty_cells,
        "saved_by_audit": False,
    }


def flatten_python(text: str) -> str:
    text = re.sub(r"from drying_model\.([a-zA-Z0-9_]+) import", r"from \1 import", text)
    text = re.sub(
        r"import drying_model\.([a-zA-Z0-9_]+) as ([a-zA-Z0-9_]+)",
        r"import \1 as \2",
        text,
    )
    text = text.replace(
        'Path(__file__).resolve().parents[1] / "scripts" / ',
        'Path(__file__).resolve().parent / ',
    ).replace(
        'Path(__file__).resolve().parents[1] / "tools" / ',
        'Path(__file__).resolve().parent / ',
    )
    replacements = {
        "data/raw/attachment1.csv": "attachment1.csv",
        "data/raw/attachment2.csv": "attachment2.csv",
        "tmp/problem1_result.json": "problem1_result.json",
        "tmp/problem2_result.json": "problem2_result.json",
        "tmp/problem3_result.json": "problem3_result.json",
        "tmp/problem4_result.json": "problem4_result.json",
        "reports/figures": ".",
        "reports/data/problem4_production_diagnostics.json": "problem4_production_diagnostics.json",
        "reports/data/problem4_review_experiments.json": "problem4_review_experiments.json",
        "reports/data/final_all_questions_audit.json": "final_all_questions_audit.json",
        "reports/problem4_independent_validation.json": "problem4_independent_validation.json",
        "outputs/result1.xlsx": "result1.xlsx",
        "outputs/result2.xlsx": "result2.xlsx",
        "outputs/result2_full_process.npz": "result2_full_process.npz",
        "outputs/result2_full_process.xlsx": "result2_full_process.xlsx",
        "outputs/result3.xlsx": "result3.xlsx",
        "outputs/result4.xlsx": "result4.xlsx",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    return text


def copy_file(source: Path, destination_name: str | None = None) -> Path:
    destination = STAGING / (destination_name or source.name)
    assert not destination.exists(), f"Duplicate filename: {destination.name}"
    shutil.copy2(source, destination)
    return destination


def build() -> None:
    for path in (STAGING, BACKUP):
        if path.exists():
            assert path.resolve().parent == OUT.resolve()
            shutil.rmtree(path)
    STAGING.mkdir()

    copy_file(SOURCE / "requirements.txt")
    for filename in ("attachment1.csv", "attachment2.csv"):
        copy_file(SOURCE / "data" / "raw" / filename)
    for filename in ("result1.xlsx", "result2.xlsx", "result3.xlsx", "result4.xlsx"):
        copy_file(SOURCE / "outputs" / filename)
    copy_file(SOURCE / "outputs" / "result2_full_process.npz")
    for source in sorted((SOURCE / "reports" / "figures").glob("*.png")):
        copy_file(source)
    for filename in (
        "final_all_questions_audit.json",
        "problem2_full_process_audit.json",
        "problem4_production_diagnostics.json",
        "problem4_review_experiments.json",
        "problem4_workbook_audit.json",
    ):
        copy_file(SOURCE / "reports" / "data" / filename)
    copy_file(SOURCE / "reports" / "problem4_independent_validation.json")
    copy_file(OUT / "AI工具使用详情.pdf", "AI 工具使用详情.pdf")

    python_sources = [
        *(ROOT / "src" / "drying_model" / name for name in (
            "__init__.py", "problem1.py", "problem2.py", "problem3.py", "problem4.py",
            "problem4_independent.py",
        )),
        *(ROOT / "scripts" / name for name in (
            "run_problem1.py", "run_problem2.py", "run_problem3.py", "run_problem4.py",
            "verify_problem4.py", "verify_problem4_independent.py",
        )),
        *(ROOT / "tests" / name for name in (
            "test_final_review_regressions.py", "test_problem1.py", "test_problem2.py",
            "test_problem3.py", "test_problem4.py", "test_problem4_independent.py",
            "test_result4_workbook.py",
        )),
        *(ROOT / "tools" / name for name in (
            "audit_all_results.py", "build_result4.py", "compress_problem2_full_process.py",
            "extend_problem2_delivery.py", "rebuild_delivery_workbooks.py",
            "restore_full_result2.py",
        )),
    ]
    for source in python_sources:
        destination = STAGING / source.name
        assert not destination.exists(), f"Duplicate filename: {destination.name}"
        destination.write_text(flatten_python(source.read_text(encoding="utf-8")), encoding="utf-8")

    readme = """# A题支撑材料

本压缩包采用单层结构，所有文件均直接位于根目录，不含子文件夹。`result1.xlsx`、`result2.xlsx`、`result3.xlsx`、`result4.xlsx` 已填写计算结果，沿用题目指定文件名和现有表格格式；打包过程仅逐字节复制，没有用表格软件重新保存。

`attachment1.csv` 和 `attachment2.csv` 为数值输入。支撑材料共含 25 份 Python 源码：四问模型与独立验证模块、运行入口、测试程序以及工作簿构建和审计工具。`problem1.py` 至 `problem4.py` 为附录中展示的核心模型，完整源码均保存在本目录。安装依赖后可先运行全部测试，再依次运行四问入口：

```powershell
python -m pip install -r requirements.txt
python -m pytest -q -p no:cacheprovider
python run_problem1.py
python run_problem2.py
python run_problem3.py
python run_problem4.py
```

四个结果工作簿是已审定交付值，运行脚本不会覆盖它们。第二问全程逐秒结果以 `result2_full_process.npz` 保存；运行 `python restore_full_result2.py` 可在根目录生成 `result2_full_process.xlsx`，脚本发现同名文件时会拒绝覆盖。

`flat_support_audit.json` 记录四个结果工作簿的 SHA-256、工作表尺寸、非空单元格数和模板样式核验结果。第四问中药材收缩后位于实体外的固定位置按题意留空，不属于漏填。

`AI 工具使用详情.pdf` 为按竞赛规定提供的独立披露文件。论文参考文献仅收录赛题与学术文献，不将 AI 工具或技能项目列入参考文献。
"""
    (STAGING / "README_复现说明.md").write_text(readme, encoding="utf-8")

    workbook_reports = []
    for index in range(1, 5):
        filename = f"result{index}.xlsx"
        source_result = SOURCE / "outputs" / filename
        staged_result = STAGING / filename
        assert sha256(source_result) == sha256(staged_result)
        workbook_reports.append(workbook_audit(TEMPLATE_DIR / filename, staged_result))

    # Import every flattened model module before replacing the previous tree.
    sys.dont_write_bytecode = True
    sys.path.insert(0, str(STAGING))
    for module_name in ("problem1", "problem2", "problem3", "problem4", "problem4_independent"):
        __import__(module_name)
    sys.path.pop(0)
    bytecode_cache = STAGING / "__pycache__"
    if bytecode_cache.exists():
        shutil.rmtree(bytecode_cache)
    test_environment = os.environ.copy()
    test_environment["PYTHONDONTWRITEBYTECODE"] = "1"
    completed = subprocess.run(
        [sys.executable, "-m", "pytest", "-q", "-p", "no:cacheprovider"],
        cwd=STAGING,
        env=test_environment,
        check=True,
        capture_output=True,
        text=True,
    )
    assert "95 passed" in completed.stdout
    assert not any(path.is_dir() for path in STAGING.iterdir())

    audit = {
        "structure": "single-level",
        "subdirectory_count": 0,
        "result_workbooks": workbook_reports,
        "result_filenames": [f"result{i}.xlsx" for i in range(1, 5)],
        "workbooks_resaved": False,
        "python_source_files": len(python_sources),
        "flat_package_tests": "95 passed",
    }
    (STAGING / "flat_support_audit.json").write_text(
        json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    assert not any(path.is_dir() for path in STAGING.iterdir())

    if TARGET.exists():
        assert TARGET.resolve().parent == OUT.resolve()
        TARGET.rename(BACKUP)
    STAGING.rename(TARGET)
    try:
        files = sorted(path for path in TARGET.iterdir() if path.is_file())
        with zipfile.ZipFile(OUT / "A题支撑材料.zip", "w", zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
            for path in files:
                archive.write(path, path.name)
        with zipfile.ZipFile(OUT / "A题支撑材料.zip") as archive:
            assert archive.testzip() is None
            assert all("/" not in name and "\\" not in name for name in archive.namelist())
            assert {path.name for path in files} == set(archive.namelist())
    except Exception:
        if TARGET.exists():
            shutil.rmtree(TARGET)
        if BACKUP.exists():
            BACKUP.rename(TARGET)
        raise
    if BACKUP.exists():
        shutil.rmtree(BACKUP)
    print(json.dumps(audit, ensure_ascii=False))


if __name__ == "__main__":
    build()
