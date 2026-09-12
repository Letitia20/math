"""Add a full-process one-second workbook without changing the original results."""
from pathlib import Path
import json
import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from problem1 import load_chamber_history_csv, sample_solution
from problem2 import solve_problem2
from problem3 import extend_chamber_history_to_plateau


def main():
    out = Path('deliverables/supporting_materials')
    (out / 'outputs').mkdir(parents=True, exist_ok=True)
    Path('tmp/paper').mkdir(parents=True, exist_ok=True)
    end = int(json.loads(Path('problem3_result.json').read_text(encoding='utf-8'))['strict_completion_time_s'])
    history = extend_chamber_history_to_plateau(load_chamber_history_csv('attachment1.csv'), end)
    radii = np.arange(21) / 10
    for n in (2560, 5120):
        cache = Path(f'tmp/paper/q2full_{n}.npz')
        if cache.exists():
            continue
        T = np.empty((end, 21)); C = np.empty_like(T)
        ti = ci = None
        for start in range(0, end, 600):
            stop = min(start + 600, end)
            times = np.arange(start + 1, stop + 1, dtype=float)
            sol = solve_problem2(history, radial_intervals=n, output_times_s=times,
                relative_tolerance=2e-10, max_step_s=2 if start < 14400 else 60,
                initial_time_s=float(start) if start else None,
                initial_temperature_c=ti, initial_moisture_concentration=ci)
            sample = sample_solution(sol, radii)
            T[start:stop] = sample.temperature_c; C[start:stop] = sample.moisture_concentration
            ti = sol.temperature_c[-1].copy(); ci = sol.moisture_concentration[-1].copy()
            if start % 21600 == 0:
                print(f'grid={n}, time={stop}/{end}', flush=True)
        np.savez_compressed(cache, temperature=T, moisture=C)
    coarse = np.load('tmp/paper/q2full_2560.npz')
    fine = np.load('tmp/paper/q2full_5120.npz')
    values = {key: (4*fine[key]-coarse[key])/3 for key in ('temperature','moisture')}
    p2 = json.loads(Path('problem2_result.json').read_text(encoding='utf-8'))
    p3 = json.loads(Path('problem3_result.json').read_text(encoding='utf-8'))
    audit = {'time_start_s':1,'time_end_s':end,'radial_intervals':[2560,5120], 'output_interval_s':1}
    for field, key in [('temperature','temperature_c'),('moisture','moisture_concentration')]:
        difference = np.abs(np.round(values[field][:10800],4)-np.array(p2[key]))
        audit[field+'_first_3h_max_difference'] = float(difference.max())
        if difference.max() > 1e-8:
            raise ValueError('First three hours differ from audited original')
    minute_rows = [(int(t)-1, i) for i,t in enumerate(p3['time_s']) if float(t).is_integer() and int(t)%60 == 0]
    diff = np.array([np.round(values['moisture'][a],4)-np.array(p3['moisture_concentration'][b]) for a,b in minute_rows])
    audit['q3_minute_max_rounded_difference'] = float(np.abs(diff).max())
    audit['strict_final_maximum_unrounded'] = float(values['moisture'][-1].max())
    if audit['strict_final_maximum_unrounded'] >= .15 or np.abs(diff).max() > .00010001:
        raise ValueError('Full-process cross-check failed')
    book = Workbook(write_only=True)
    for name, field in [('温度','temperature'),('水分浓度','moisture')]:
        sheet = book.create_sheet(name)
        sheet.append(['时间\\到药材中心的距离']+list(radii))
        for i, row in enumerate(np.round(values[field],4)):
            cells = [i+1]
            for v in row:
                cell = WriteOnlyCell(sheet, float(v)); cell.number_format='0.0000'; cells.append(cell)
            sheet.append(cells)
    dest = out/'result2.xlsx'
    book.save(dest)
    check = load_workbook(dest,read_only=True,data_only=True)
    checked=0
    for sheet, field in zip(check.worksheets,('temperature','moisture')):
        for i,row in enumerate(sheet.iter_rows(min_row=2,values_only=True)):
            if row[0] != i+1 or not np.allclose(row[1:],np.round(values[field][i],4),rtol=0,atol=1e-12):
                raise ValueError('Workbook readback mismatch')
            checked+=len(row)
        if i+1 != end: raise ValueError('Missing output rows')
    check.close()
    audit['readback_checked_cells']=checked
    audit['workbook_bytes']=dest.stat().st_size
    audit['status']='passed'
    (out/'reports/data').mkdir(parents=True,exist_ok=True)
    (out/'reports/data/problem2_full_process_audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(audit),flush=True)


if __name__=='__main__': main()
