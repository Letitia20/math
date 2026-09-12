"""Build the four workbooks from the independently reproduced JSON payloads."""
from pathlib import Path
import json
from openpyxl import Workbook

def main():
    out=Path('outputs/recomputed'); out.mkdir(parents=True,exist_ok=True)
    for number in (1,2,3,4):
        payload=json.loads(Path(f'tmp/problem{number}_result.json').read_text(encoding='utf-8'))
        book=Workbook(); book.remove(book.active)
        for name,key in ([('温度','temperature_c'),('水分浓度','moisture_concentration')] if number<3 else [('Sheet1','moisture_concentration')]):
            ws=book.create_sheet(name)
            ws.append(['时间\\到药材中心的距离']+payload['radius_cm']+(['药材表面'] if number==4 else []))
            for i,(t,row) in enumerate(zip(payload['time_s'],payload[key],strict=True)):
                ws.append([t]+row+([payload['surface_moisture_concentration'][i]] if number==4 else []))
            for row in ws.iter_rows(min_row=2,min_col=2):
                for c in row: c.number_format='0.0000'
        dest=out/f'result{number}.xlsx'
        if dest.exists(): raise FileExistsError(dest)
        book.save(dest)
        print(dest)

if __name__=='__main__': main()
