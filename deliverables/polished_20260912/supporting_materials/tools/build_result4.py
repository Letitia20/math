"""Fill and audit the official Problem 4 workbook with numeric results."""

from __future__ import annotations

import argparse
import csv
import json
import math
from bisect import bisect_left
from copy import copy
from pathlib import Path
from typing import Any

import openpyxl

DEFAULT_RADIUS_HISTORY = Path("data/raw/attachment2.csv")
HEADER = "时间\\到药材中心的距离"
FIXED_RADII_CM = [index / 10.0 for index in range(21)]
# Match the solver's 1e-12 m geometric tolerance, expressed here in cm.
RADIUS_TOLERANCE_CM = 1.0e-10
CRITICAL_MOISTURE_TOLERANCE = 1.0e-7


def _finite_number(value: Any, label: str) -> float:
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        raise ValueError(f"{label} must be a finite numeric value")
    return float(value)


def _radius_history(path: Path) -> tuple[list[float], list[float]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    times = [float(row["time_s"]) for row in rows]
    radii = [float(row["radius_cm"]) for row in rows]
    if len(times) < 2 or any(not math.isfinite(value) for value in times + radii):
        raise ValueError("Radius history must contain at least two finite observations")
    if any(right <= left for left, right in zip(times, times[1:])):
        raise ValueError("Radius-history times must be strictly increasing")
    if any(value <= 0.0 for value in radii):
        raise ValueError("Radius-history values must be positive")
    return times, radii


def _radius_at(query: float, times: list[float], radii: list[float]) -> float:
    """Interpolate measured radii, holding the last radius after observation."""
    if query < times[0]:
        raise ValueError("Output time precedes the measured radius-history horizon")
    if query >= times[-1]:
        return radii[-1]
    index = bisect_left(times, query)
    if times[index] == query:
        return radii[index]
    fraction = (query - times[index - 1]) / (times[index] - times[index - 1])
    return radii[index - 1] + fraction * (radii[index] - radii[index - 1])


def validate_payload(
    payload: dict[str, Any],
    radius_history: Path = DEFAULT_RADIUS_HISTORY,
) -> dict[str, Any]:
    """Validate time, geometry, and unrounded endpoint evidence.

    Time identities use exact equality: relative tolerances can incorrectly
    merge a minute sample and a nearby critical sample at large times.
    Geometry is recomputed from the observations, not the payload blank mask.
    """
    times = payload["time_s"]
    radii = payload["radius_cm"]
    matrix = payload["moisture_concentration"]
    surface = payload["surface_moisture_concentration"]
    if radii != FIXED_RADII_CM or any(isinstance(value, bool) for value in radii):
        raise ValueError("Fixed radius headers must be 0, 0.1, ..., 2.0 cm")
    if not times or len(times) != len(matrix) or len(times) != len(surface):
        raise ValueError("Payload dimensions do not match the result4 template")
    if any(len(row) != 21 for row in matrix):
        raise ValueError("Moisture matrix has an invalid shape")
    numeric_times = [_finite_number(value, "Output time") for value in times]
    if any(right <= left for left, right in zip(numeric_times, numeric_times[1:])):
        raise ValueError("Output times must be strictly increasing")
    critical = _finite_number(payload["drying_time_s"], "Critical time")
    strict = _finite_number(payload["strict_completion_time_s"], "Strict completion time")
    if critical <= 0.0 or strict != 60.0 * (math.floor(critical / 60.0) + 1):
        raise ValueError("Strict completion must be the first minute strictly after the critical time")
    if numeric_times.count(critical) != 1:
        raise ValueError("Exactly one critical time row is required")
    expected_times = sorted(set([float(value) for value in range(60, int(strict) + 1, 60)] + [critical]))
    if numeric_times != expected_times:
        raise ValueError("Time rows must contain every 60 s sample and exactly one critical row")
    history_times, history_radii = _radius_history(Path(radius_history))
    blank_count = 0
    for index, time in enumerate(numeric_times):
        actual_radius = _radius_at(time, history_times, history_radii)
        for radius, value in zip(FIXED_RADII_CM, matrix[index]):
            if radius > actual_radius + RADIUS_TOLERANCE_CM:
                if value is not None:
                    raise ValueError(f"Outside-radius cell must be blank at time {time}, radius {radius}")
                blank_count += 1
            else:
                moisture = _finite_number(value, f"Inside-radius moisture at time {time}, radius {radius}")
                if moisture < 0.0:
                    raise ValueError("Moisture values must not be negative")
        if _finite_number(surface[index], f"Surface moisture at time {time}") < 0.0:
            raise ValueError("Surface moisture values must not be negative")
    threshold = _finite_number(payload["threshold"], "Moisture threshold")
    if threshold != 0.15:
        raise ValueError("Problem 4 requires the 0.15 kg/kg moisture threshold")
    critical_maximum = _finite_number(payload["critical_maximum_moisture_unrounded"], "Unrounded critical maximum")
    strict_maximum = _finite_number(payload["strict_completion_maximum_moisture_unrounded"], "Unrounded strict maximum")
    if abs(critical_maximum - threshold) > CRITICAL_MOISTURE_TOLERANCE:
        raise ValueError("Unrounded critical maximum does not meet the threshold tolerance")
    if not 0.0 <= strict_maximum < threshold:
        raise ValueError("Unrounded strict completion maximum must be strictly below the threshold")
    for time, maximum in ((critical, critical_maximum), (strict, strict_maximum)):
        index = numeric_times.index(time)
        displayed_maximum = max([value for value in matrix[index] if value is not None] + [surface[index]])
        if abs(displayed_maximum - round(maximum, 4)) > 1.0e-12:
            raise ValueError("Rounded endpoint values disagree with the unrounded maximum")
    return {
        "data_rows": len(times),
        "rows": len(times) + 1,
        "columns": 23,
        "blank_outside_cells": blank_count,
        "radius_last_observation_time_s": history_times[-1],
        "radius_extension": "hold_last_value",
        "rows_after_radius_observations": sum(time > history_times[-1] for time in numeric_times),
        "critical_time_s": critical,
        "strict_completion_time_s": strict,
        "critical_maximum_moisture_unrounded": critical_maximum,
        "strict_completion_maximum_moisture_unrounded": strict_maximum,
    }


def audit_result4(
    workbook_path: Path,
    payload: dict[str, Any] | Path,
    radius_history: Path = DEFAULT_RADIUS_HISTORY,
) -> dict[str, Any]:
    """Reopen the workbook and compare every cell with the verified payload.

    No workbook metadata is removed. Unexpected comments, hidden content,
    formulas, errors, extra sheets, and changed numeric values fail the audit.
    """
    if not isinstance(payload, dict):
        payload = json.loads(Path(payload).read_text(encoding="utf-8"))
    summary = validate_payload(payload, radius_history)
    workbook = openpyxl.load_workbook(workbook_path, data_only=False)
    try:
        for sheet in workbook.worksheets:
            if sheet.sheet_state != "visible":
                raise ValueError(f"Hidden worksheet: {sheet.title}")
            if sheet.sheet_format.zeroHeight or any(dimension.hidden for dimension in sheet.row_dimensions.values()):
                raise ValueError(f"Hidden rows in {sheet.title}")
            if any(dimension.hidden for dimension in sheet.column_dimensions.values()):
                raise ValueError(f"Hidden columns in {sheet.title}")
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.comment is not None:
                        raise ValueError(f"Unexpected comment in {sheet.title}!{cell.coordinate}")
                    if cell.data_type in ("f", "e"):
                        raise ValueError(f"Unexpected formula or error in {sheet.title}!{cell.coordinate}")
        if workbook.sheetnames != ["Sheet1"]:
            raise ValueError("Result4 must contain only the official Sheet1")
        sheet = workbook["Sheet1"]
        if (sheet.max_row, sheet.max_column) != (summary["rows"], summary["columns"]):
            raise ValueError("Workbook dimensions do not match the validated payload")
        expected = [[HEADER, *FIXED_RADII_CM, "药材表面"]]
        expected.extend(
            [time, *row, surface]
            for time, row, surface in zip(payload["time_s"], payload["moisture_concentration"], payload["surface_moisture_concentration"])
        )
        for actual_row, expected_row in zip(sheet.iter_rows(), expected):
            for cell, value in zip(actual_row, expected_row):
                if cell.value != value or isinstance(cell.value, bool):
                    raise ValueError(f"Workbook/payload mismatch at {cell.coordinate}: {cell.value!r} != {value!r}")
        return {"status": "passed", **summary, "checked_cells": summary["rows"] * summary["columns"]}
    finally:
        workbook.close()


def build_result4(
    template: Path,
    payload: dict[str, Any] | Path,
    output: Path,
    radius_history: Path = DEFAULT_RADIUS_HISTORY,
) -> dict[str, Any]:
    """Populate the template, preserving its numeric-output conventions."""
    if not isinstance(payload, dict):
        payload = json.loads(Path(payload).read_text(encoding="utf-8"))
    summary = validate_payload(payload, radius_history)
    times = payload["time_s"]
    radii = payload["radius_cm"]
    matrix = payload["moisture_concentration"]
    surface = payload["surface_moisture_concentration"]
    workbook = openpyxl.load_workbook(template)
    if "Sheet1" not in workbook.sheetnames:
        raise ValueError("Official result4 template must contain Sheet1")
    sheet = workbook["Sheet1"]
    sheet.cell(1, 1).value = HEADER
    for col, radius in enumerate(radii, start=2):
        sheet.cell(1, col).value = radius
    sheet.cell(1, len(radii) + 2).value = "药材表面"
    required_rows = summary["rows"]
    if sheet.max_row > required_rows:
        sheet.delete_rows(required_rows + 1, sheet.max_row - required_rows)
    for row_index, time in enumerate(times, start=2):
        for col_index, value in enumerate([time, *matrix[row_index - 2], surface[row_index - 2]], start=1):
            sheet.cell(row_index, col_index).value = value

    sheet.column_dimensions["A"].width = 24
    for col in range(2, len(radii) + 3):
        sheet.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 11
    for row in sheet.iter_rows(min_row=1, max_row=required_rows, min_col=1, max_col=len(radii) + 2):
        for cell in row:
            alignment = copy(cell.alignment)
            alignment.horizontal = "center"
            alignment.vertical = "center"
            cell.alignment = alignment
            if cell.row == 1:
                font = copy(cell.font)
                font.name = "宋体"
                font.size = 10
                cell.font = font
            elif cell.column == 1:
                cell.number_format = "0.######"
            else:
                cell.number_format = "0.0000"
    sheet.freeze_panes = "B2"
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    workbook.close()
    return audit_result4(output, payload, radius_history)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("template", type=Path)
    parser.add_argument("payload", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--radius-history", type=Path, default=DEFAULT_RADIUS_HISTORY)
    parser.add_argument("--audit-only", action="store_true", help="Audit an existing output; do not write a workbook")
    parser.add_argument("--audit-output", type=Path, help="Write the successful audit summary as JSON")
    args = parser.parse_args()
    if args.audit_output is not None and args.audit_output.resolve() in {
        path.resolve() for path in (args.template, args.payload, args.output, args.radius_history)
    }:
        parser.error("--audit-output must be distinct from the template, payload, workbook, and radius history")
    if args.audit_only:
        summary = audit_result4(args.output, args.payload, args.radius_history)
    else:
        summary = build_result4(args.template, args.payload, args.output, args.radius_history)
    report = {"workbook": str(args.output), **summary}
    if args.audit_output is not None:
        args.audit_output.parent.mkdir(parents=True, exist_ok=True)
        args.audit_output.write_text(json.dumps(report, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False))


if __name__ == "__main__":
    main()
