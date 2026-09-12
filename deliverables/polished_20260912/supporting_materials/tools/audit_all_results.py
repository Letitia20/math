"""Read-only audit of original inputs, four result workbooks and computed payloads."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from pathlib import Path

import numpy as np
import openpyxl

from build_result4 import audit_result4

HEADER = "时间\\到药材中心的距离"
RADII = [i / 10 for i in range(21)]


def file_hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def check_original_inputs(attachment_dir: Path) -> dict:
    results = {}
    for number, columns in ((1, 3), (2, 2)):
        source = attachment_dir / f"附件{number}.xlsx"
        book = openpyxl.load_workbook(source, read_only=True, data_only=True)
        try:
            original = np.asarray([row[:columns] for row in book.worksheets[0].iter_rows(min_row=2, values_only=True)
                                   if row[0] is not None], dtype=float)
        finally:
            book.close()
        csv_path = Path(f"data/raw/attachment{number}.csv")
        with csv_path.open(encoding="utf-8-sig", newline="") as stream:
            reader = csv.reader(stream)
            next(reader)
            exported = np.asarray(list(reader), dtype=float)
        if original.shape != exported.shape or not np.array_equal(original, exported):
            raise ValueError(f"Attachment {number} CSV differs from original workbook")
        results[str(number)] = {"rows": len(original), "columns": columns,
                                "exact_numeric_match": True, "original_sha256": file_hash(source),
                                "csv_sha256": file_hash(csv_path), "first_row": original[0].tolist(),
                                "last_row": original[-1].tolist()}
    return results


def audit_fixed_workbook(path: Path, payload: dict, number: int) -> dict:
    times = payload["time_s"]
    if payload["radius_cm"] != RADII:
        raise ValueError(f"Question {number}: invalid fixed radii")
    if number in (1, 2):
        expected_count = 1800 if number == 1 else 10800
        if times != list(range(1, expected_count + 1)):
            raise ValueError(f"Question {number}: incomplete one-second output")
        fields = [("温度", "temperature_c"), ("水分浓度", "moisture_concentration")]
    else:
        critical = payload["drying_time_s"]
        strict = payload["strict_completion_time_s"]
        expected_strict = (np.floor(critical / 60) + 1) * 60
        expected_times = sorted(set(list(range(60, int(expected_strict) + 1, 60)) + [critical]))
        if strict != expected_strict or times != expected_times:
            raise ValueError("Question 3: missing minute, critical or first strict row")
        if abs(payload["critical_maximum_moisture_unrounded"] - 0.15) > 1e-8:
            raise ValueError("Question 3: invalid unrounded critical maximum")
        if not 0 <= payload["strict_completion_maximum_moisture_unrounded"] < 0.15:
            raise ValueError("Question 3: invalid unrounded strict maximum")
        fields = [("Sheet1", "moisture_concentration")]
    book = openpyxl.load_workbook(path, data_only=False)
    summaries = []
    mismatches = []
    mismatch_count = 0
    try:
        if book.sheetnames != [name for name, _ in fields]:
            raise ValueError(f"Question {number}: invalid sheets")
        for name, field in fields:
            sheet = book[name]
            if sheet.sheet_state != "visible" or sheet.sheet_format.zeroHeight:
                raise ValueError(f"Hidden content in {path}:{name}")
            if any(dim.hidden for dim in [*sheet.row_dimensions.values(), *sheet.column_dimensions.values()]):
                raise ValueError(f"Hidden rows or columns in {path}:{name}")
            expected_shape = (len(times) + 1, 22)
            if (sheet.max_row, sheet.max_column) != expected_shape:
                raise ValueError(f"Question {number}: invalid dimensions {sheet.max_row, sheet.max_column}")
            data = np.asarray(payload[field], dtype=float)
            if data.shape != (len(times), 21) or not np.isfinite(data).all():
                raise ValueError(f"Question {number}: invalid payload field")
            if field == "moisture_concentration" and np.min(data) < 0:
                raise ValueError("Negative moisture in payload")
            if not np.array_equal(data, np.round(data, 4)):
                raise ValueError("Output concentrations/temperatures must be rounded to four decimals")
            expected = [[HEADER, *RADII], *[[time, *row] for time, row in zip(times, payload[field])]]
            for actual_row, expected_row in zip(sheet.iter_rows(), expected):
                for cell, value in zip(actual_row, expected_row):
                    if cell.data_type in ("f", "e") or cell.comment is not None:
                        raise ValueError(f"Unexpected formula, error or comment at {name}!{cell.coordinate}")
                    if cell.value != value or isinstance(cell.value, bool):
                        mismatch_count += 1
                        if len(mismatches) < 12:
                            mismatches.append({"sheet": name, "cell": cell.coordinate, "workbook": cell.value, "recomputed": value})
                    if cell.row > 1 and cell.column > 1 and cell.number_format != "0.0000":
                        raise ValueError(f"Incorrect number format at {name}!{cell.coordinate}")
            summaries.append({"name": name, "rows": sheet.max_row, "columns": sheet.max_column,
                              "checked_cells": sheet.max_row * sheet.max_column})
    finally:
        book.close()
    return {"status": "passed" if mismatch_count == 0 else "mismatch", "sheets": summaries,
            "mismatch_count": mismatch_count, "first_mismatches": mismatches,
            "workbook_sha256": file_hash(path)}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--attachment-dir", type=Path, required=True)
    parser.add_argument("--payload-dir", type=Path, default=Path("tmp"))
    parser.add_argument("--output", type=Path, default=Path("reports/data/final_all_questions_audit.json"))
    args = parser.parse_args()
    results = {"inputs": check_original_inputs(args.attachment_dir), "workbooks": {}, "numerical_evidence": {}}
    payloads = {}
    for number in range(1, 5):
        payload_path = args.payload_dir / f"problem{number}_result.json"
        payload = json.loads(payload_path.read_text(encoding="utf-8"))
        payloads[number] = payload
        workbook = Path(f"outputs/result{number}.xlsx")
        audit = audit_result4(workbook, payload) if number == 4 else audit_fixed_workbook(workbook, payload, number)
        audit["payload_sha256"] = file_hash(payload_path)
        audit["workbook_sha256"] = file_hash(workbook)
        results["workbooks"][str(number)] = audit
        excluded = {"time_s", "radius_cm", "temperature_c", "moisture_concentration", "surface_moisture_concentration"}
        results["numerical_evidence"][str(number)] = {k: v for k, v in payload.items() if k not in excluded}
        print(f"Question {number}: {audit['status']}", flush=True)
    q2 = np.asarray(payloads[2]["moisture_concentration"])[59::60]
    q3 = np.asarray(payloads[3]["moisture_concentration"])[:180]
    results["q2_q3_shared_first_three_hours"] = {
        "same_continuum_model": True, "comparison_points": int(q2.size),
        "maximum_absolute_rounded_difference": float(np.max(np.abs(q2 - q3))),
        "rounded_mismatch_count": int(np.count_nonzero(q2 != q3)),
        "note": "Q2 and Q3 use different time-integration tolerances; four-decimal boundary flips are assessed separately.",
    }
    tables = []
    for number, fields, hours in ((1, ["temperature_c", "moisture_concentration"], False),
                                  (2, ["temperature_c", "moisture_concentration"], True),
                                  (3, ["moisture_concentration"], True), (4, ["moisture_concentration"], True)):
        report = Path(f"reports/problem{number}.md").read_text(encoding="utf-8")
        payload = payloads[number]
        for table_index, field in enumerate(fields):
            table_number = (1 if number == 1 else 3 if number == 2 else 5 if number == 3 else 6) + table_index
            section = re.split(rf"\*\*表 {table_number}\s", report, maxsplit=1)[1]
            rows = []
            started = False
            for line in section.splitlines():
                if line.startswith("|"):
                    started = True
                    cells = [value.strip() for value in line.strip("|").split("|")]
                    if re.fullmatch(r"\d+(?:\.\d+)?", cells[0]):
                        rows.append(cells)
                elif started:
                    break
            for cells in rows:
                time = float(cells[0]) * (3600 if hours else 1)
                index = payload["time_s"].index(time)
                expected = [payload[field][index][i] for i in ([0, 5, 10] if number == 4 else [0, 5, 10, 15, 20])]
                if number == 4:
                    expected.append(payload["surface_moisture_concentration"][index])
                if [float(value) for value in cells[1:]] != expected:
                    raise ValueError(f"Table {table_number} differs from recomputed values at {time} s")
            expected_rows = 7 if number == 1 else 6 if number == 2 else 9 if number == 3 else 8
            if len(rows) != expected_rows:
                raise ValueError(f"Table {table_number} is incomplete")
            tables.append({"table": table_number, "checked_regular_rows": len(rows), "status": "passed"})
    results["paper_tables"] = tables
    results["status"] = "passed" if all(row["status"] == "passed" for row in results["workbooks"].values()) else "mismatch"
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(results, indent=2, ensure_ascii=False, allow_nan=False) + "\n", encoding="utf-8")
    print(json.dumps({"status": results["status"], "shared_Q2_Q3": results["q2_q3_shared_first_three_hours"]}, ensure_ascii=False))
    if results["status"] != "passed":
        raise SystemExit(1)


if __name__ == "__main__":
    main()
