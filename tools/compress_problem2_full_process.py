"""Convert the large full-process workbook to a compact compressed NPZ archive."""
from pathlib import Path
import numpy as np
from openpyxl import load_workbook

root = Path(__file__).resolve().parents[1]
src = root / 'deliverables' / 'supporting_materials' / 'outputs' / 'result2.xlsx'
dst = src.with_name('result2_full_process.npz')
wb = load_workbook(src, read_only=True, data_only=True)
arrays = {}
for ws in wb.worksheets:
    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    data = np.asarray([[np.nan if v is None else v for v in row] for row in rows], dtype=np.float64)
    arrays[ws.title] = data
    arrays[f'{ws.title}_header'] = np.asarray([str(v) for v in header])
np.savez_compressed(dst, **arrays)
print({'source_bytes': src.stat().st_size, 'npz_bytes': dst.stat().st_size,
       'sheets': list(arrays), 'rows': {k: v.shape for k, v in arrays.items() if not k.endswith('_header')}})
