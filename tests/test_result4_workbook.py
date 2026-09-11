"""Round-trip and deliberate-corruption checks for the submission workbook."""

from __future__ import annotations

import importlib.util
import json
import math
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest
from openpyxl.comments import Comment


BUILDER_PATH = Path(__file__).resolve().parents[1] / "tools" / "build_result4.py"
SPEC = importlib.util.spec_from_file_location("result4_workbook_builder", BUILDER_PATH)
builder = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(builder)


def _fixture(tmp_path: Path, critical: float = 119.9999):
    radius_path = tmp_path / "radius.csv"
    radius_path.write_text("time_s,radius_cm\n0,2\n300,0.5\n", encoding="utf-8")
    strict = 60 * (math.floor(critical / 60) + 1)
    times = sorted(set([*range(60, strict + 1, 60), critical]))
    radii = [index / 10 for index in range(21)]
    matrix = []
    for time in times:
        maximum = 0.15 if time == critical else 0.14998 if time == strict else 0.25
        actual_radius = 2 - time / 200
        matrix.append([
            round(maximum - index * 0.002, 4) if radius <= actual_radius + 1e-10 else None
            for index, radius in enumerate(radii)
        ])
    payload = {
        "time_s": times,
        "radius_cm": radii,
        "moisture_concentration": matrix,
        "surface_moisture_concentration": [0.04] * len(times),
        "drying_time_s": critical,
        "strict_completion_time_s": strict,
        "threshold": 0.15,
        "critical_maximum_moisture_unrounded": 0.15,
        "strict_completion_maximum_moisture_unrounded": 0.14998,
    }
    template = tmp_path / "template.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "Sheet1"
    workbook.active["A1"] = builder.HEADER
    workbook.active["A20"] = "placeholder"
    workbook.properties.title = "Official numeric result template"
    workbook.save(template)
    workbook.close()
    return template, payload, tmp_path / "result4.xlsx", radius_path


@pytest.mark.parametrize("critical", [119.9999, 120.0, 120.0001, 30.0])
def test_roundtrip_preserves_nearby_critical_and_minute_rows(tmp_path, critical):
    template, payload, output, radius = _fixture(tmp_path, critical)
    summary = builder.build_result4(template, payload, output, radius)
    assert summary["status"] == "passed"
    assert summary["rows"] == len(payload["time_s"]) + 1
    assert summary["columns"] == 23
    workbook = openpyxl.load_workbook(output)
    sheet = workbook["Sheet1"]
    assert [cell.value for cell in sheet["A"][1:]] == payload["time_s"]
    assert sheet["A1"].comment is None
    assert sheet["B2"].number_format == "0.0000"
    assert workbook.properties.title == "Official numeric result template"
    # Strict compliance relies on unrounded evidence: both display as 0.1500.
    assert sheet.cell(sheet.max_row, 2).value == 0.15
    assert summary["strict_completion_maximum_moisture_unrounded"] < 0.15
    workbook.close()


@pytest.mark.parametrize(
    "corruption",
    [
        "inside_blank", "outside_zero", "surface_blank", "numeric_change",
        "radius_header", "surface_header", "time_change", "missing_minute",
        "formula", "error", "comment", "hidden_row", "hidden_column",
        "hidden_sheet", "extra_sheet", "extra_column", "boolean", "hidden_by_default",
    ],
)
def test_audit_rejects_workbook_corruption(tmp_path, corruption):
    template, payload, output, radius = _fixture(tmp_path)
    builder.build_result4(template, payload, output, radius)
    workbook = openpyxl.load_workbook(output)
    sheet = workbook["Sheet1"]
    if corruption == "inside_blank":
        sheet["B2"] = None
    elif corruption == "outside_zero":
        sheet["V2"] = 0
    elif corruption == "surface_blank":
        sheet["W2"] = None
    elif corruption == "numeric_change":
        sheet["C2"] = sheet["C2"].value + 0.0001
    elif corruption == "radius_header":
        sheet["C1"] = 0.11
    elif corruption == "surface_header":
        sheet["W1"] = "2 cm"
    elif corruption == "time_change":
        sheet["A3"] = 120.0  # Replaces the nearby critical time.
    elif corruption == "missing_minute":
        sheet.delete_rows(2)
    elif corruption == "formula":
        sheet["B2"] = "=0.25"
    elif corruption == "error":
        sheet["B2"] = "#VALUE!"
    elif corruption == "comment":
        sheet["A1"].comment = Comment("Unexpected annotation", "Reviewer")
    elif corruption == "hidden_row":
        sheet.row_dimensions[2].hidden = True
    elif corruption == "hidden_column":
        sheet.column_dimensions["C"].hidden = True
    elif corruption == "hidden_by_default":
        sheet.sheet_format.zeroHeight = True
    elif corruption == "hidden_sheet":
        workbook.create_sheet("hidden").sheet_state = "hidden"
    elif corruption == "extra_sheet":
        workbook.create_sheet("extra")
    elif corruption == "extra_column":
        sheet["X1"] = "unexpected"
    elif corruption == "boolean":
        sheet["B1"] = False  # Python equality alone would equate this to 0.
    workbook.save(output)
    workbook.close()
    with pytest.raises(ValueError):
        builder.audit_result4(output, payload, radius)


@pytest.mark.parametrize(
    "corruption",
    [
        "inside_blank", "outside_zero", "surface_nan", "nonfinite_inside",
        "bad_headers", "missing_minute", "nearby_critical", "critical_error",
        "strict_not_below", "strict_rounding_mismatch", "wrong_threshold",
        "duplicate_time", "skipped_strict_minute", "wrong_geometry",
    ],
)
def test_audit_rejects_invalid_payload_even_when_workbook_could_match(tmp_path, corruption):
    _, payload, _, radius = _fixture(tmp_path)
    if corruption == "inside_blank":
        payload["moisture_concentration"][0][0] = None
    elif corruption == "outside_zero":
        payload["moisture_concentration"][0][-1] = 0
    elif corruption == "surface_nan":
        payload["surface_moisture_concentration"][0] = float("nan")
    elif corruption == "nonfinite_inside":
        payload["moisture_concentration"][0][0] = float("inf")
    elif corruption == "bad_headers":
        payload["radius_cm"][1] = 0.11
    elif corruption == "missing_minute":
        for key in ("time_s", "moisture_concentration", "surface_moisture_concentration"):
            payload[key].pop(0)
    elif corruption == "nearby_critical":
        payload["time_s"][1] = 119.9998
    elif corruption == "critical_error":
        payload["critical_maximum_moisture_unrounded"] = 0.1500002
    elif corruption == "strict_not_below":
        payload["strict_completion_maximum_moisture_unrounded"] = 0.15
    elif corruption == "strict_rounding_mismatch":
        payload["strict_completion_maximum_moisture_unrounded"] = 0.14
    elif corruption == "wrong_threshold":
        payload["threshold"] = 0.2
    elif corruption == "duplicate_time":
        payload["time_s"][-1] = payload["time_s"][-2]
    elif corruption == "skipped_strict_minute":
        payload["strict_completion_time_s"] += 60
    elif corruption == "wrong_geometry":
        radius.write_text("time_s,radius_cm\n0,2\n300,2\n", encoding="utf-8")
    with pytest.raises(ValueError):
        builder.validate_payload(payload, radius)


def test_original_cli_and_readonly_audit_mode(tmp_path):
    template, payload, output, radius = _fixture(tmp_path)
    payload_path = tmp_path / "payload.json"
    payload_path.write_text(json.dumps(payload), encoding="utf-8")
    command = [sys.executable, str(BUILDER_PATH), str(template), str(payload_path), str(output), "--radius-history", str(radius)]
    created = subprocess.run(command, capture_output=True, text=True, check=True)
    assert json.loads(created.stdout)["status"] == "passed"
    original_bytes = output.read_bytes()
    checked = subprocess.run([*command, "--audit-only"], capture_output=True, text=True, check=True)
    assert json.loads(checked.stdout)["status"] == "passed"
    assert output.read_bytes() == original_bytes


def test_radius_extension_holds_the_last_observed_value(tmp_path):
    template, payload, output, radius = _fixture(tmp_path)
    radius.write_text("time_s,radius_cm\n0,2\n60,1.7\n", encoding="utf-8")
    for time, row in zip(payload["time_s"], payload["moisture_concentration"]):
        maximum = 0.15 if time == payload["drying_time_s"] else 0.14998 if time == payload["strict_completion_time_s"] else 0.25
        for index, position in enumerate(payload["radius_cm"]):
            row[index] = round(maximum - index * 0.002, 4) if position <= 1.7 else None
    summary = builder.build_result4(template, payload, output, radius)
    assert summary["status"] == "passed"
    assert summary["rows_after_radius_observations"] == 2
    assert summary["radius_extension"] == "hold_last_value"
    assert builder._radius_at(90, [0, 60], [2, 1.7]) == 1.7
    assert builder._radius_at(60, [0, 60], [2, 1.7]) == 1.7
    with pytest.raises(ValueError, match="precedes"):
        builder._radius_at(-1, [0, 60], [2, 1.7])
    # A tail row must use the held radius, so its 1.7 cm cell cannot be blank.
    payload["moisture_concentration"][-1][17] = None
    with pytest.raises(ValueError, match="Inside-radius"):
        builder.validate_payload(payload, radius)


@pytest.mark.parametrize("audit_only", [False, True])
def test_cli_can_archive_audit_json(tmp_path, audit_only):
    template, payload, output, radius = _fixture(tmp_path)
    payload_path = tmp_path / "payload.json"
    payload_path.write_text(json.dumps(payload), encoding="utf-8")
    archive = tmp_path / "reports" / "data" / "workbook_audit.json"
    command = [sys.executable, str(BUILDER_PATH), str(template), str(payload_path), str(output), "--radius-history", str(radius), "--audit-output", str(archive)]
    if audit_only:
        builder.build_result4(template, payload, output, radius)
        previous_workbook = output.read_bytes()
        command.append("--audit-only")
    result = subprocess.run(command, capture_output=True, text=True, check=True)
    assert json.loads(archive.read_text(encoding="utf-8")) == json.loads(result.stdout)
    assert json.loads(result.stdout)["status"] == "passed"
    if audit_only:
        assert output.read_bytes() == previous_workbook


def test_audit_archive_cannot_overwrite_the_workbook(tmp_path):
    template, payload, output, radius = _fixture(tmp_path)
    payload_path = tmp_path / "payload.json"
    payload_path.write_text(json.dumps(payload), encoding="utf-8")
    builder.build_result4(template, payload, output, radius)
    previous_workbook = output.read_bytes()
    command = [sys.executable, str(BUILDER_PATH), str(template), str(payload_path), str(output), "--radius-history", str(radius), "--audit-output", str(output), "--audit-only"]
    result = subprocess.run(command, capture_output=True, text=True)
    assert result.returncode != 0
    assert "must be distinct" in result.stderr
    assert output.read_bytes() == previous_workbook
